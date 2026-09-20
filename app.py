"""
Resume Bot — Streamlit app
---------------------------
Generate a resume from a prompt, OR upload an existing resume + context
to have it rewritten. Export as PDF, DOCX, or TEX.

Setup:
    Set GROQ_API_KEY as an environment variable (or in a .env file next
    to this script). No key is ever entered in the UI.

Run:
    streamlit run app.py
"""

import os
import io
import re
import base64
import streamlit as st
from groq import Groq

# ---- Optional parsing / rendering libs ----
import pdfplumber
from docx import Document
from docx.shared import Pt
from fpdf import FPDF, XPos, YPos
from fpdf.enums import WrapMode

# Load a local .env file if python-dotenv is available. Not required —
# GROQ_API_KEY can just as easily be set in the real environment.
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass


# =========================================================
# FIXED CONFIG — no key or model selection lives in the UI
# =========================================================

GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "")
GROQ_MODEL = "openai/gpt-oss-120b"  # hardcode here if you want a different model

ACCENT = "#FF4B4B"


# =========================================================
# LLM CALL
# =========================================================

SYSTEM_PROMPT = """You are an expert resume writer.
Given context about a candidate (and optionally their existing resume text),
produce a polished, well-structured resume.

Output ONLY in the following Markdown structure, nothing else,
no commentary, no code fences, no backticks:

# Full Name
City, Country | email | phone | linkedin/portfolio (if known, else omit)

## Summary
2-3 sentence professional summary.

## Experience
### Job Title — Company (Start – End)
- Bullet point achievement
- Bullet point achievement

## Education
### Degree — Institution (Year)

## Skills
- Skill, Skill, Skill

If the candidate's context includes any of the following, add them as
additional sections using this exact heading text, placed after Skills,
in this order — only include a section if there is real content for it:

## Projects
### Project Name (Tech used, if given)
- Bullet describing what it does / your contribution

## Certifications
- Certification Name — Issuer (Year)

## Languages
- Language (Proficiency level)

## Awards
- Award Name — Issuer (Year)

## Volunteer Experience
### Role — Organization (Start – End)
- Bullet describing the work

Only include sections that have real content. Do not invent facts that
contradict what the user gave you, but you may phrase/organize freely.
Use a plain "-" for every bullet point, never "•" or other symbols.
"""


def call_llm(existing_resume_text: str, user_context: str) -> str:
    client = Groq(api_key=GROQ_API_KEY)

    if existing_resume_text.strip():
        user_msg = (
            f"Here is the candidate's existing resume text:\n\n"
            f"{existing_resume_text}\n\n"
            f"Here is the new context / instructions for how to change it:\n\n"
            f"{user_context}\n\n"
            f"Rewrite the resume accordingly."
        )
    else:
        user_msg = (
            f"Here is context about the candidate. Generate a resume from scratch:\n\n"
            f"{user_context}"
        )

    resp = client.chat.completions.create(
        model=GROQ_MODEL,
        max_tokens=3000,
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": user_msg},
        ],
    )
    return resp.choices[0].message.content


def clean_llm_output(text: str) -> str:
    """Defensive cleanup for LLM output that doesn't perfectly follow the
    requested format — strips stray code fences and normalizes bullet
    glyphs to plain markdown '-' so the parser below doesn't choke."""
    if not text:
        return text
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\n?", "", text)
        text = re.sub(r"\n?```$", "", text)
    text = re.sub(r"^[ \t]*[•●▪◦‣][ \t]*", "- ", text, flags=re.MULTILINE)
    text = re.sub(r"^[ \t]*\*[ \t]+", "- ", text, flags=re.MULTILINE)
    return text.strip()


# =========================================================
# FILE PARSING (uploaded resume -> plain text)
# =========================================================

def parse_uploaded_resume(uploaded_file) -> str:
    if uploaded_file is None:
        return ""

    name = uploaded_file.name.lower()

    if name.endswith(".pdf"):
        text = []
        with pdfplumber.open(uploaded_file) as pdf:
            for page in pdf.pages:
                text.append(page.extract_text() or "")
        return "\n".join(text)

    elif name.endswith(".docx"):
        doc = Document(uploaded_file)
        return "\n".join(p.text for p in doc.paragraphs)

    elif name.endswith(".tex") or name.endswith(".txt"):
        return uploaded_file.read().decode("utf-8", errors="ignore")

    else:
        st.warning(f"Unsupported file type: {uploaded_file.name}")
        return ""


# =========================================================
# MARKDOWN -> STRUCTURED SECTIONS (shared by all renderers)
# =========================================================

def parse_markdown_resume(md_text: str):
    """
    Small parser for the fixed structure the LLM is asked to output.
    Returns a dict: {name, contact, sections: [{heading, entries}]}
    """
    md_text = clean_llm_output(md_text)
    lines = [l.rstrip() for l in md_text.strip().splitlines() if l.strip() != ""]
    data = {"name": "", "contact": "", "sections": []}

    i = 0
    if lines and lines[0].startswith("# "):
        data["name"] = lines[0][2:].strip()
        i = 1
    if i < len(lines) and not lines[i].startswith("#"):
        data["contact"] = lines[i].strip()
        i += 1

    current_section = None
    current_sub = None

    for line in lines[i:]:
        if line.startswith("## "):
            current_section = {"heading": line[3:].strip(), "entries": []}
            data["sections"].append(current_section)
            current_sub = None
        elif line.startswith("### "):
            current_sub = {"title": line[4:].strip(), "bullets": []}
            if current_section is not None:
                current_section["entries"].append(current_sub)
        elif line.startswith("- "):
            bullet = line[2:].strip()
            if current_sub is not None:
                current_sub["bullets"].append(bullet)
            elif current_section is not None:
                current_section["entries"].append({"title": None, "bullets": [bullet]})
        else:
            if current_section is not None:
                current_section["entries"].append({"title": None, "bullets": [line.strip()]})

    return data


# =========================================================
# RENDERERS
# =========================================================

def render_tex(data) -> str:
    def esc(s):
        return (s.replace("&", r"\&").replace("%", r"\%").replace("#", r"\#")
                 .replace("_", r"\_"))

    lines = [
        r"\documentclass[11pt]{article}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage[margin=1in]{geometry}",
        r"\usepackage{enumitem}",
        r"\pagestyle{empty}",
        r"\begin{document}",
        r"\begin{center}",
        r"{\Large \textbf{" + esc(data["name"]) + r"}} \\",
        esc(data["contact"]) + r" \\",
        r"\end{center}",
    ]
    for section in data["sections"]:
        lines.append(r"\section*{" + esc(section["heading"]) + "}")
        for entry in section["entries"]:
            if entry["title"]:
                lines.append(r"\textbf{" + esc(entry["title"]) + r"} \\")
            if entry["bullets"]:
                lines.append(r"\begin{itemize}[leftmargin=*]")
                for b in entry["bullets"]:
                    lines.append(r"\item " + esc(b))
                lines.append(r"\end{itemize}")
    lines.append(r"\end{document}")
    return "\n".join(lines)


def render_docx(data) -> bytes:
    doc = Document()

    title = doc.add_paragraph()
    run = title.add_run(data["name"])
    run.bold = True
    run.font.size = Pt(18)

    doc.add_paragraph(data["contact"])

    for section in data["sections"]:
        doc.add_heading(section["heading"], level=2)
        for entry in section["entries"]:
            if entry["title"]:
                p = doc.add_paragraph()
                p.add_run(entry["title"]).bold = True
            for b in entry["bullets"]:
                doc.add_paragraph(b, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_PDF_CHAR_MAP = {
    "\u2022": "-", "\u2023": "-", "\u25E6": "-",
    "\u2013": "-", "\u2014": "--",
    "\u2018": "'", "\u2019": "'",
    "\u201C": '"', "\u201D": '"',
    "\u2026": "...", "\u00A0": " ",
}


def sanitize_for_pdf(text: str) -> str:
    """Normalize 'smart' punctuation the LLM likes to produce to plain
    ASCII equivalents. With real embedded TTF fonts (Unicode-capable) this
    is now just cosmetic consistency, not a crash-prevention hack."""
    if not text:
        return text
    for bad, good in _PDF_CHAR_MAP.items():
        text = text.replace(bad, good)
    return text


# ---------------------------------------------------------
# PDF templates
# ---------------------------------------------------------
# All three are strictly single-column with standard section headings and
# plain "-" bullets — no tables, columns, icons, or text boxes — because
# those are exactly the layout patterns that trip up ATS resume parsers.
# The only things that differ are font, color accent, and spacing.

FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")

PDF_TEMPLATES = {
    "Minimal": {
        "family": "LiberationSans",  # metric-compatible with Arial
        "regular": os.path.join(FONT_DIR, "LiberationSans-Regular.ttf"),
        "bold": os.path.join(FONT_DIR, "LiberationSans-Bold.ttf"),
        "accent": (35, 35, 35),
        "rule": (205, 205, 205),
        "name_size": 19, "header_size": 11.5, "body_size": 10.5,
        "uppercase_headers": False,
    },
    "Modern": {
        "family": "DejaVuSans",
        "regular": os.path.join(FONT_DIR, "DejaVuSans.ttf"),
        "bold": os.path.join(FONT_DIR, "DejaVuSans-Bold.ttf"),
        "accent": (0, 88, 156),
        "rule": (0, 88, 156),
        "name_size": 20, "header_size": 11.5, "body_size": 10.2,
        "uppercase_headers": True,
    },
    "Executive": {
        "family": "LiberationSerif",  # metric-compatible with Times New Roman
        "regular": os.path.join(FONT_DIR, "LiberationSerif-Regular.ttf"),
        "bold": os.path.join(FONT_DIR, "LiberationSerif-Bold.ttf"),
        "accent": (89, 22, 22),
        "rule": (89, 22, 22),
        "name_size": 21, "header_size": 12, "body_size": 11,
        "uppercase_headers": True,
    },
}


def render_pdf(data, template_name: str = "Minimal") -> bytes:
    tpl = PDF_TEMPLATES.get(template_name, PDF_TEMPLATES["Minimal"])
    fam = tpl["family"]

    pdf = FPDF()
    pdf.add_font(fam, "", tpl["regular"])
    pdf.add_font(fam, "B", tpl["bold"])
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_margins(18, 16, 18)
    pdf.set_x(pdf.l_margin)

    def reset_x():
        pdf.set_x(pdf.l_margin)  # multi_cell(w=0) leaves x near the right margin,
                                  # which would starve the next call's width and
                                  # hang WrapMode.CHAR if not reset every time

    def hrule(color, thickness=0.5):
        pdf.set_draw_color(*color)
        pdf.set_line_width(thickness)
        pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())

    # --- Header: name + contact ---
    pdf.set_font(fam, "B", tpl["name_size"])
    pdf.set_text_color(*tpl["accent"])
    pdf.cell(0, tpl["name_size"] * 0.5, sanitize_for_pdf(data["name"]),
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.set_font(fam, "", tpl["body_size"])
    pdf.set_text_color(70, 70, 70)
    pdf.cell(0, 6.5, sanitize_for_pdf(data["contact"]), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(1.5)
    hrule(tpl["rule"], thickness=0.7)
    pdf.ln(4)

    # --- Sections ---
    for section in data["sections"]:
        heading = section["heading"].upper() if tpl["uppercase_headers"] else section["heading"]
        pdf.set_font(fam, "B", tpl["header_size"])
        pdf.set_text_color(*tpl["accent"])
        reset_x()
        pdf.cell(0, 7, sanitize_for_pdf(heading), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        hrule(tpl["rule"], thickness=0.3)
        pdf.ln(2.5)

        pdf.set_text_color(30, 30, 30)
        for entry in section["entries"]:
            if entry["title"]:
                pdf.set_font(fam, "B", tpl["body_size"])
                reset_x()
                pdf.multi_cell(0, 5.8, sanitize_for_pdf(entry["title"]), wrapmode=WrapMode.CHAR)
                pdf.ln(0.5)
            pdf.set_font(fam, "", tpl["body_size"])
            for b in entry["bullets"]:
                reset_x()
                pdf.multi_cell(0, 5.8, f"-  {sanitize_for_pdf(b)}", wrapmode=WrapMode.CHAR)
            pdf.ln(1.5)
        pdf.ln(2.5)

    return bytes(pdf.output())


# =========================================================
# STREAMLIT UI
# =========================================================

st.set_page_config(page_title="Resume Bot", page_icon="📄", layout="wide")

st.markdown(f"""
<style>
    .block-container {{ padding-top: 2rem; max-width: 1100px; }}
    h1, h2, h3 {{ letter-spacing: -0.01em; }}
    div[data-testid="stForm"], .resume-card {{
        border-radius: 14px;
    }}
    .upload-chip {{
        display: inline-block; padding: 4px 12px; border-radius: 999px;
        background: rgba(255,75,75,0.12); color: {ACCENT};
        font-size: 0.85rem; font-weight: 600; margin-top: 6px;
    }}
    div.stButton > button[kind="primary"] {{
        background-color: {ACCENT}; border-color: {ACCENT};
        font-weight: 600; height: 3em;
    }}
    .resume-card h1 {{ margin-bottom: 0.1rem; font-size: 1.8rem; }}
    .resume-card h2 {{
        border-bottom: 2px solid {ACCENT}; padding-bottom: 4px;
        margin-top: 1.2rem; font-size: 1.15rem; text-transform: uppercase;
        letter-spacing: 0.05em;
    }}
    .resume-card h3 {{ font-size: 1rem; margin-bottom: 0.2rem; }}
</style>
""", unsafe_allow_html=True)

st.title("📄 Resume Bot")
st.caption("Generate a new resume, or upload one and tell it what to change.")

if not GROQ_API_KEY:
    st.error(
        "GROQ_API_KEY is not set. Set it as an environment variable "
        "(or in a .env file next to app.py) and restart the app."
    )
    st.stop()

show_debug = st.query_params.get("debug") == "true"

# ---------------------------------------------------------
# Mode selection
# ---------------------------------------------------------
mode = st.radio(
    "What do you want to do?",
    ["✨ Create from scratch", "✏️ Edit an existing resume"],
    horizontal=True,
)
editing_mode = mode.startswith("✏️")

uploaded_file = None
if editing_mode:
    uploaded_file = st.file_uploader(
        "Upload your existing resume",
        type=["pdf", "docx", "tex", "txt"],
    )
    if uploaded_file is not None:
        st.markdown(f'<span class="upload-chip">✓ {uploaded_file.name} loaded</span>',
                    unsafe_allow_html=True)

context_label = (
    "What should change?" if editing_mode
    else "Tell me about the candidate"
)
context_placeholder = (
    "e.g. 'Tailor this for a data analyst role, emphasize SQL and "
    "dashboarding work, shorten the summary.'"
    if editing_mode else
    "Role, years of experience, key skills, education, notable "
    "achievements — anything relevant."
)

context = st.text_area(context_label, height=180, placeholder=context_placeholder)

# ---------------------------------------------------------
# Optional extra sections — a lightweight checklist instead of a full
# structured form. Ticking a box reveals a small guided text area with an
# example, so people discover these exist without having to know to type
# them into the main box. Still free text under the hood — the LLM does
# the organizing — just with a nudge toward what's worth mentioning.
# ---------------------------------------------------------
EXTRA_SECTIONS = {
    "Projects": "e.g. 'Budget tracker app — React + Firebase, used by 200+ users, "
                "cut manual expense entry time in half.'",
    "Certifications": "e.g. 'AWS Certified Solutions Architect (2024), "
                       "Google Data Analytics Certificate (2023)'",
    "Languages": "e.g. 'English (fluent), Tamil (native), Spanish (conversational)'",
    "Awards": "e.g. 'Best Intern Project Award — Acme Corp (2023), "
              "Dean's List all semesters'",
    "Volunteer Experience": "e.g. 'Tutor — local NGO (2021–present), "
                             "taught basic coding to high schoolers on weekends'",
}

with st.expander("➕ Add more to your resume (projects, certifications, languages...)"):
    chosen_extras = st.multiselect(
        "What else do you want included?",
        list(EXTRA_SECTIONS.keys()),
    )
    extra_section_text = {}
    for section_name in chosen_extras:
        extra_section_text[section_name] = st.text_area(
            section_name, placeholder=EXTRA_SECTIONS[section_name], height=80,
            key=f"extra_{section_name}",
        )

col1, col2 = st.columns([2, 1])
with col1:
    generate_clicked = st.button("Generate résumé", type="primary", use_container_width=True)
with col2:
    out_format = st.radio("Format", ["PDF", "DOCX", "TEX"], horizontal=True, label_visibility="collapsed")

pdf_template = st.radio(
    "PDF style", list(PDF_TEMPLATES.keys()), horizontal=True,
    help="Applies to the PDF export. All templates are single-column and ATS-safe "
         "(no tables, columns, icons, or text boxes — just font, color, and spacing).",
) if out_format == "PDF" else "Minimal"

# ---------------------------------------------------------
# Generation
# ---------------------------------------------------------
if generate_clicked:
    if editing_mode and uploaded_file is None:
        st.error("Upload a resume to edit, or switch to 'Create from scratch'.")
    elif not context.strip():
        st.error("Add some context or instructions.")
    else:
        with st.spinner("Reading input..."):
            existing_text = parse_uploaded_resume(uploaded_file) if editing_mode else ""

        # Fold any filled-in guided sections into the context as clearly
        # labeled blocks, so the LLM has real content to place under the
        # matching "## Projects" / "## Certifications" / etc. heading.
        full_context = context.strip()
        for section_name, text in extra_section_text.items():
            if text and text.strip():
                full_context += f"\n\n{section_name}:\n{text.strip()}"

        with st.spinner("Writing your resume..."):
            try:
                md_result = call_llm(existing_text, full_context)
            except Exception as e:
                st.error(f"Generation failed: {e}")
                st.stop()

        # Seed the editor widget's own state key BEFORE it is instantiated
        # below — this is what actually fixes the "can't regenerate" bug
        # from writing to the same key a widget also owns.
        st.session_state["md_editor"] = clean_llm_output(md_result)
        st.toast("Resume generated!", icon="✅")

# ---------------------------------------------------------
# Preview + export
# ---------------------------------------------------------
if "md_editor" in st.session_state and st.session_state["md_editor"].strip():
    current_md = st.session_state["md_editor"]
    data = parse_markdown_resume(current_md)

    if not data["sections"]:
        st.warning(
            "Couldn't detect any sections in the generated text — check the "
            "raw content below and edit if needed before exporting."
        )

    st.subheader("Preview")
    with st.container(border=True):
        st.markdown(f'<div class="resume-card">', unsafe_allow_html=True)
        st.markdown(current_md)
        st.markdown('</div>', unsafe_allow_html=True)

    with st.expander("Edit raw content"):
        st.text_area("Markdown source", key="md_editor", height=300,
                      label_visibility="collapsed")

    if show_debug:
        with st.expander("Debug: parsed structure"):
            st.json(data)

    st.subheader("Export")
    try:
        if out_format == "TEX":
            content = render_tex(data)
            st.download_button("Download .tex", content, file_name="resume.tex",
                                mime="text/x-tex", use_container_width=True)
        elif out_format == "DOCX":
            content = render_docx(data)
            st.download_button("Download .docx", content, file_name="resume.docx",
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                use_container_width=True)
        elif out_format == "PDF":
            content = render_pdf(data, template_name=pdf_template)
            st.download_button("Download .pdf", content, file_name="resume.pdf",
                                mime="application/pdf", use_container_width=True)

            # Real preview of the exact bytes being downloaded, not an
            # approximation — the browser's own PDF viewer renders this,
            # so no extra dependency or conversion step is needed.
            b64_pdf = base64.b64encode(content).decode("utf-8")
            st.markdown(
                f'<iframe src="data:application/pdf;base64,{b64_pdf}" '
                f'width="100%" height="800" style="border:1px solid #444; border-radius:8px;" '
                f'type="application/pdf"></iframe>',
                unsafe_allow_html=True,
            )
            st.caption(
                "If the preview doesn't render (some mobile browsers don't support "
                "inline PDFs), the download button above always has the real file."
            )
    except Exception as e:
        st.error("Rendering the export file failed. Details below:")
        st.exception(e)

    if st.button("🔄 Start over"):
        del st.session_state["md_editor"]
        st.rerun()